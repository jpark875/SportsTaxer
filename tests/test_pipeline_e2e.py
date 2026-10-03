"""Whole pipeline on a synthetic recording, with the model replaced by recorded answers."""

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sportstaxer.cli import main
from sportstaxer.config import Config
from sportstaxer.extract import ModelField, ModelResponse, ModelRow, ReplayBackend, plan_tiles
from sportstaxer.stitch import load_canvas_manifest
from sportstaxer.synthetic import make_capture

from .conftest import encode

ROW_HEIGHT = 96
CONFIG = """\
extract:
  tile_height: 600
  tile_overlap: 150
  money_passes: 1
work_dir: work
out_dir: out
adapters_dir: adapters
"""

ADAPTER = """\
name: plainbook
display_name: Plain Book
payout_includes_stake: false
date_formats: ["%Y-%m-%d"]
"""


def field(value: str) -> ModelField:
    return ModelField(value=value, confidence=0.97)


def money(value: float) -> str:
    return f"${value:,.2f}"


def net_of(row) -> Decimal:
    if row.result == "Lost":
        return -Decimal(str(row.stake))
    if row.result == "Push":
        return Decimal("0")
    return Decimal(str(row.payout))


@pytest.fixture
def workspace(tmp_path):
    (tmp_path / "adapters").mkdir()
    (tmp_path / "adapters" / "plainbook.yaml").write_text(ADAPTER)
    (tmp_path / "sportstaxer.yaml").write_text(CONFIG)
    capture = make_capture(row_count=14, frame_height=640, step=300)
    video = encode(capture.frames, tmp_path / "recording.mp4")
    return tmp_path, capture, video


def record_responses(capture, canvas_dir: Path, directory: Path) -> None:
    """What a perfect model would say about each tile: every row that fits inside it."""
    config = Config(extract={"tile_height": 600, "tile_overlap": 150})
    manifest = load_canvas_manifest(canvas_dir)
    directory.mkdir()
    for tile in plan_tiles(manifest, canvas_dir, config):
        rows = []
        for index, source in enumerate(capture.rows):
            top, bottom = index * ROW_HEIGHT, (index + 1) * ROW_HEIGHT
            if top < tile.y_offset or bottom > tile.y_offset + 600:
                continue
            rows.append(
                ModelRow(
                    date=field(source.date),
                    description=field(source.description),
                    bet_type=ModelField(value=None, confidence=0.0),
                    stake=field(money(source.stake)),
                    payout=field(money(source.payout)),
                    result=field(source.result),
                    balance=ModelField(value=None, confidence=0.0),
                    y_top=top - tile.y_offset,
                    y_bottom=bottom - tile.y_offset - 1,
                )
            )
        path = directory / ReplayBackend.filename(tile)
        path.write_text(ModelResponse(rows=rows).model_dump_json())


def test_recording_to_reconciled_workbook(workspace, monkeypatch, capsys):
    root, capture, video = workspace
    monkeypatch.chdir(root)

    assert main(["frames", str(video)]) == 0
    assert main(["stitch", "recording", "--book", "plainbook"]) == 0
    record_responses(capture, root / "work" / "recording" / "canvas", root / "responses")

    opening = Decimal("1000.00")
    closing = opening + sum(net_of(r) for r in capture.rows)
    code = main(
        [
            "extract",
            "recording",
            "--book",
            "plainbook",
            "--backend",
            "replay",
            "--responses",
            str(root / "responses"),
        ]
    )
    assert code == 0
    code = main(
        [
            "ledger",
            "recording",
            "--book",
            "plainbook",
            "--opening",
            str(opening),
            "--closing",
            str(closing),
        ]
    )
    assert code == 0
    assert "reconciliation: balanced" in capsys.readouterr().out

    workbook = load_workbook(root / "out" / "recording.xlsx")
    sheet = workbook["Ledger"]
    assert sheet.max_row == len(capture.rows) + 1
    assert sheet.cell(row=2, column=2).value == capture.rows[0].description
    assert workbook["Summary"]["B2"].value == "balanced"


def test_wrong_closing_balance_fails_the_run(workspace, monkeypatch, capsys):
    root, capture, video = workspace
    monkeypatch.chdir(root)
    main(["frames", str(video)])
    main(["stitch", "recording", "--book", "plainbook"])
    record_responses(capture, root / "work" / "recording" / "canvas", root / "responses")
    main(
        [
            "extract",
            "recording",
            "--book",
            "plainbook",
            "--backend",
            "replay",
            "--responses",
            str(root / "responses"),
        ]
    )

    code = main(["ledger", "recording", "--book", "plainbook", "--opening", "0", "--closing", "1"])

    assert code == 3
    assert "reconciliation FAILED" in capsys.readouterr().err
    assert (root / "out" / "recording.xlsx").exists()


def test_run_command_chains_every_stage(workspace, monkeypatch, capsys):
    root, capture, video = workspace
    monkeypatch.chdir(root)
    main(["frames", str(video)])
    main(["stitch", "recording", "--book", "plainbook"])
    record_responses(capture, root / "work" / "recording" / "canvas", root / "responses")

    code = main(
        [
            "run",
            str(video),
            "--run-id",
            "again",
            "--book",
            "plainbook",
            "--force",
            "--backend",
            "replay",
            "--responses",
            str(root / "responses"),
        ]
    )

    assert code == 0
    assert (root / "out" / "again.xlsx").exists()
    assert "reconciliation: unverified" in capsys.readouterr().out


def test_replay_without_responses_is_an_error(workspace, monkeypatch, capsys):
    root, _, video = workspace
    monkeypatch.chdir(root)
    main(["frames", str(video)])
    main(["stitch", "recording", "--book", "plainbook"])

    code = main(["extract", "recording", "--book", "plainbook", "--backend", "replay"])

    assert code == 1
    assert "--responses" in capsys.readouterr().err
