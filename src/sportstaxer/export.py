"""Stage 6: write the ledger, a summary and a review list to xlsx."""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sportstaxer.parse import LedgerRow, Result
from sportstaxer.reconcile import Reconciliation

MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00"
HEADER_FONT = Font(bold=True)
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")

LEDGER_COLUMNS = [
    ("Date", 12),
    ("Description", 44),
    ("Type", 16),
    ("Result", 11),
    ("Stake", 12),
    ("Payout", 12),
    ("Net", 12),
    ("Balance", 12),
    ("Review", 48),
]


def _money(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _header(sheet, titles: list[tuple[str, int]]) -> None:
    for column, (title, width) in enumerate(titles, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"


def _ledger_sheet(workbook: Workbook, rows: list[LedgerRow]) -> None:
    sheet = workbook.active
    sheet.title = "Ledger"
    _header(sheet, LEDGER_COLUMNS)
    for index, row in enumerate(rows, start=2):
        values = [
            row.date,
            row.description,
            row.bet_type,
            row.result.value,
            _money(row.stake),
            _money(row.payout),
            _money(row.net),
            _money(row.balance),
            "; ".join(row.flags),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            if column == 1 and row.date:
                cell.number_format = "yyyy-mm-dd"
            elif 5 <= column <= 8:
                cell.number_format = MONEY_FORMAT
            if row.needs_review:
                cell.fill = FLAG_FILL
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(LEDGER_COLUMNS))}{max(len(rows) + 1, 2)}"


def _summary_sheet(workbook: Workbook, rows: list[LedgerRow], recon: Reconciliation) -> None:
    sheet = workbook.create_sheet("Summary")
    sheet.column_dimensions["A"].width = 26
    for letter in "BCDE":
        sheet.column_dimensions[letter].width = 14

    sheet["A1"] = "Reconciliation"
    sheet["A1"].font = HEADER_FONT
    facts = [
        ("Status", recon.status),
        ("Opening balance", _money(recon.opening)),
        ("Sum of net", _money(sum((r.net for r in rows if r.net is not None), Decimal("0")))),
        ("Expected closing", _money(recon.expected_closing)),
        ("Reported closing", _money(recon.closing)),
        ("Difference", _money(recon.difference)),
        ("Rows needing review", recon.flagged),
    ]
    line = 2
    for label, value in facts:
        sheet.cell(row=line, column=1, value=label)
        cell = sheet.cell(row=line, column=2, value=value)
        if isinstance(value, float):
            cell.number_format = MONEY_FORMAT
        line += 1
    for note in recon.notes:
        sheet.cell(row=line, column=1, value=note)
        line += 1

    line += 1
    sheet.cell(row=line, column=1, value="By month").font = HEADER_FONT
    line += 1
    for column, title in enumerate(("Month", "Bets", "Staked", "Net"), start=1):
        sheet.cell(row=line, column=column, value=title).font = HEADER_FONT
    line += 1

    months: dict[str, list[LedgerRow]] = defaultdict(list)
    for row in rows:
        if row.result in (Result.DEPOSIT, Result.WITHDRAWAL):
            continue
        months[row.date.strftime("%Y-%m") if row.date else "undated"].append(row)
    first = line
    for month in sorted(months):
        bucket = months[month]
        sheet.cell(row=line, column=1, value=month)
        sheet.cell(row=line, column=2, value=len(bucket))
        staked = sum((r.stake or Decimal("0") for r in bucket), Decimal("0"))
        net = sum((r.net or Decimal("0") for r in bucket), Decimal("0"))
        sheet.cell(row=line, column=3, value=_money(staked)).number_format = MONEY_FORMAT
        sheet.cell(row=line, column=4, value=_money(net)).number_format = MONEY_FORMAT
        line += 1
    if months:
        sheet.cell(row=line, column=1, value="Total").font = HEADER_FONT
        for column in (2, 3, 4):
            letter = get_column_letter(column)
            cell = sheet.cell(
                row=line, column=column, value=f"=SUM({letter}{first}:{letter}{line - 1})"
            )
            cell.font = HEADER_FONT
            if column > 2:
                cell.number_format = MONEY_FORMAT


def _review_sheet(workbook: Workbook, rows: list[LedgerRow]) -> None:
    sheet = workbook.create_sheet("Review")
    _header(sheet, [("Ledger row", 12), ("Description", 44), ("Reason", 80)])
    line = 2
    for index, row in enumerate(rows, start=2):
        for flag in row.flags:
            sheet.cell(row=line, column=1, value=index)
            sheet.cell(row=line, column=2, value=row.description)
            cell = sheet.cell(row=line, column=3, value=flag)
            cell.alignment = Alignment(wrap_text=True)
            line += 1


def export_xlsx(rows: list[LedgerRow], recon: Reconciliation, path: Path) -> Path:
    """Write the workbook. `Ledger row` on the Review sheet is the sheet row number."""
    workbook = Workbook()
    _ledger_sheet(workbook, rows)
    _summary_sheet(workbook, rows, recon)
    _review_sheet(workbook, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
